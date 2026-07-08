"""Unit tests for Word report generation."""
import os
import tempfile

import pytest
from docx import Document

from engine.docx_report import build_docx, _template_path


@pytest.fixture
def sample_report_data():
    return {
        "project": {
            "name_code": "TEST_PROJECT",
            "mode": "AUTO",
            "fuel": "Diesel",
            "gears": "8AT",
            "odriv_milestone": "MS3",
            "version": "29",
            "area": "EU",
            "target_vehicle": "REF",
        },
        "global": {
            "driv": {
                "verdict": "Medium Risk",
                "verdict_pred": "Low Risk",
                "index": 72.5,
                "rate_low": 0.12,
            },
            "dyn": {
                "verdict": "Low Risk",
                "verdict_pred": "Low Risk",
                "index": 81.0,
                "rate_low": 0.05,
            },
        },
        "sdv_results": [
            {
                "name": "Drive Away Creep Eng On",
                "order": 1,
                "n_events": 5,
                "driv": {
                    "index": 70.2,
                    "target_index": 75.0,
                    "status": {"1": "RED", "2": "GREEN", "3": "GREEN"},
                    "status_pred": {"1": "ORANGE", "2": "GREEN", "3": "GREEN"},
                    "counts": {"RED_P1": 1, "GREEN_P1": 4},
                    "lowest_event": "SubEvent-1",
                },
                "dyn": {
                    "index": 78.0,
                    "target_index": 80.0,
                    "status": {"1": "GREEN", "2": "GREEN", "3": "GREEN"},
                    "status_pred": {"1": "GREEN", "2": "GREEN", "3": "GREEN"},
                    "counts": {"GREEN_P1": 5},
                    "lowest_event": None,
                },
            }
        ],
        "charts": {"Drive Away Creep Eng On": {"png": None}},
        "events_by_sdv": {
            "Drive Away Creep Eng On": [
                {
                    "channels": [{"name": "Sub Event Name", "value": "E1"}],
                    "driv": {
                        "criticity": 1,
                        "priority": 1,
                        "color": "RED",
                        "indice_occ": -2.5,
                    },
                }
            ]
        },
        "doc_versions": [{"name": "ODRIV", "version": "v29 test"}],
    }


def test_template_exists():
    assert os.path.isfile(_template_path())


def test_build_docx_creates_valid_file(sample_report_data):
    fd, out = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        build_docx(
            sample_report_data,
            out,
            report_fields={"sender_name": "Unit Test", "domain": "DRIVABILITY-DYNAMISM"},
        )
        assert os.path.getsize(out) > 10000
        doc = Document(out)
        assert len(doc.tables) >= 6
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Drive Away Creep Eng On".upper() in text.upper()
        assert doc.tables[0].rows[2].cells[1].text.strip() == "TEST_PROJECT"
    finally:
        if os.path.exists(out):
            os.unlink(out)
