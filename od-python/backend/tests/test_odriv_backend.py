"""End-to-end backend tests for ODRIV FastAPI port."""
import io
import os
import pytest
import requests

BASE_URL = os.environ['REACT_APP_BACKEND_URL'].rstrip('/') if os.environ.get('REACT_APP_BACKEND_URL') \
    else open('/app/frontend/.env').read().split('REACT_APP_BACKEND_URL=')[1].split('\n')[0].strip()
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def session():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


# ------------------------------------------------------------- state / config
class TestStateAndConfig:
    def test_state_ok(self, session):
        r = session.get(f"{API}/state", timeout=30)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "version" in data and "Python" in data["version"]
        assert "sheets" in data and "total_events" in data

    def test_config_sections_present(self, session):
        r = session.get(f"{API}/config", timeout=30)
        assert r.status_code == 200
        cfg = r.json()
        for key in ["catalog", "definitions", "structure", "targets",
                    "settings_blocks", "target_vehicle", "chart_params",
                    "criticity", "thresholds"]:
            assert key in cfg, f"missing config section: {key}"

    def test_config_get_section(self, session):
        r = session.get(f"{API}/config/catalog", timeout=15)
        assert r.status_code == 200
        cat = r.json()
        assert isinstance(cat, list) and len(cat) >= 60  # 64 SDVs

    def test_unknown_config_section_404(self, session):
        r = session.get(f"{API}/config/__nope__", timeout=10)
        assert r.status_code == 404


# ------------------------------------------------------------- project lifecycle
class TestProjectPipeline:
    PROJECT = {
        "name_code": "TEST_PYTEST_PROJECT",
        "mode": "AUTO",
        "fuel": "Diesel",
        "gears": "8AT",
        "software_milestone": "MS3",
        "priority": "P1",
        "version": "4.6",
        "odriv_milestone": "MS3",
        "area": "EU",
        "target_vehicle": "REF",
        "number_of_gears": 8,
    }

    def test_01_new_project(self, session):
        r = session.post(f"{API}/project/new", json=self.PROJECT, timeout=15)
        assert r.status_code == 200, r.text
        p = r.json()
        assert p["name_code"] == self.PROJECT["name_code"]
        assert p["fuel"] == "Diesel"
        assert "id" in p

    def test_02_update_project(self, session):
        r = session.put(f"{API}/project", json={"area": "NA"}, timeout=15)
        assert r.status_code == 200
        p = r.json()
        assert p["area"] == "NA"
        # other fields preserved
        assert p["name_code"] == self.PROJECT["name_code"]

    def test_03_import_demo(self, session):
        r = session.post(f"{API}/import/demo", timeout=60)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["imported"] > 400, f"too few events: {data}"
        assert data["classified"] == data["imported"], "all demo events should classify"
        assert len(data["per_sdv"]) >= 50, f"only {len(data['per_sdv'])} SDVs populated"

    def test_04_state_reflects_events(self, session):
        r = session.get(f"{API}/state", timeout=15)
        data = r.json()
        assert data["total_events"] > 400
        assert data["project"]["name_code"] == self.PROJECT["name_code"]
        assert len(data["sheets"]) >= 40

    def test_05_calculate_rating(self, session):
        r = session.post(f"{API}/rating/calculate", timeout=120)
        assert r.status_code == 200, r.text
        data = r.json()
        glob = data["global"]
        assert "driv" in glob and "dyn" in glob
        for k in ["driv", "dyn"]:
            assert "index" in glob[k]
            assert "rate_low" in glob[k]
            assert glob[k]["verdict"] in {"Low Risk", "Medium Risk", "High Risk"}
        assert data["sdv_count"] >= 50

    def test_06_get_rating(self, session):
        r = session.get(f"{API}/rating", timeout=15)
        assert r.status_code == 200
        data = r.json()
        assert data["global"] is not None
        assert len(data["rows"]) >= 50
        # rows have expected structure
        row = data["rows"][0]
        assert "name" in row and "driv" in row
        if row.get("driv"):
            assert "index" in row["driv"]

    def test_07_sdv_detail(self, session):
        r = session.get(f"{API}/rating", timeout=15)
        rows = r.json()["rows"]
        first = rows[0]["name"]
        r2 = session.get(f"{API}/sdv/{first}", timeout=15)
        assert r2.status_code == 200, r2.text
        d = r2.json()
        assert d["name"]
        assert "events" in d and len(d["events"]) > 0
        assert "structure" in d and "charts" in d and "targets" in d

    def test_08_sdv_detail_unknown(self, session):
        r = session.get(f"{API}/sdv/__doesnotexist__", timeout=10)
        assert r.status_code == 404

    def test_09_events_listing(self, session):
        r = session.get(f"{API}/events?limit=10", timeout=15)
        assert r.status_code == 200
        data = r.json()
        assert data["total"] > 400
        assert len(data["events"]) == 10
        assert "id" in data["events"][0]

    def test_10_event_update_and_delete(self, session):
        r = session.get(f"{API}/events?limit=1", timeout=10)
        ev = r.json()["events"][0]
        eid = ev["id"]
        new_chans = dict(ev["channels"])
        # add a custom field
        new_chans["TEST_FIELD"] = {"min": 1, "max": 2, "mean": 1.5, "value": 1.5}
        r2 = session.put(f"{API}/events/{eid}", json={"channels": new_chans}, timeout=15)
        assert r2.status_code == 200
        updated = r2.json()
        assert "TEST_FIELD" in updated["channels"]
        # delete
        r3 = session.delete(f"{API}/events/{eid}", timeout=10)
        assert r3.status_code == 200
        # verify gone
        r4 = session.delete(f"{API}/events/{eid}", timeout=10)
        assert r4.status_code == 404

    def test_11_set_as_target(self, session):
        # New behaviour: "Set as target" reads + scores the SELECTED target
        # vehicle(s) recorded in New Project. With no target selected it errors;
        # once a target is recorded it populates the comparison column(s).
        r = session.post(f"{API}/rating/set-as-target", timeout=20)
        assert r.status_code == 400  # no target vehicle selected yet
        # record a reference target, then set-as-target should populate it
        session.post(f"{API}/targets/pending", json={
            "primary": None,
            "extras": [{"source": "ref", "id": None,
                        "label": "24M_3.6_JL_PS-A_V222_0223"}]}, timeout=20)
        r = session.post(f"{API}/rating/set-as-target", timeout=60)
        assert r.status_code == 200, r.text
        assert r.json()["ok"] is True
        rt = session.get(f"{API}/rating", timeout=20).json()
        assert "24M_3.6_JL_PS-A_V222_0223" in rt.get("comparisons", {})

    def test_12_download_sample_xlsx(self, session):
        r = session.get(f"{API}/import/sample", timeout=60)
        assert r.status_code == 200
        ctype = r.headers.get("content-type", "")
        assert "spreadsheetml" in ctype or "officedocument" in ctype
        assert len(r.content) > 5000
        assert r.content[:2] == b"PK"  # xlsx is a zip

    def test_13_reimport_via_file_endpoint(self, session):
        # download sample, then re-upload
        s2 = requests.Session()  # no Content-Type
        d = s2.get(f"{API}/import/sample", timeout=60)
        # need to wipe events first, but new project deletes all - skip wipe
        # Just verify upload works
        files = {"file": ("sample.xlsx", io.BytesIO(d.content),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/import/file", files=files, timeout=60)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["imported"] > 400
        assert data["classified"] == data["imported"]

    def test_14_report_pptx(self, session):
        # Need fresh rating since reimport added more events
        rc = session.post(f"{API}/rating/calculate", timeout=120)
        assert rc.status_code == 200
        r = session.post(f"{API}/report/pptx",
                         json={"report_fields": {}},
                         timeout=180)
        assert r.status_code == 200, r.text
        assert len(r.content) > 5000
        assert r.content[:2] == b"PK"  # pptx is a zip

    def test_15_report_pdf(self, session):
        r = session.post(f"{API}/report/pdf", json={}, timeout=180)
        assert r.status_code == 200, r.text
        assert len(r.content) > 1000
        assert r.content[:4] == b"%PDF"

    def test_15b_report_docx(self, session):
        r = session.post(
            f"{API}/report/docx",
            json={"report_fields": {"name": "Test User"}},
            timeout=180,
        )
        assert r.status_code == 200, r.text
        assert len(r.content) > 5000
        assert r.content[:2] == b"PK"

    def test_16_report_bad_fmt(self, session):
        r = session.post(f"{API}/report/xlsx", json={}, timeout=10)
        assert r.status_code == 400

    def test_17_logs(self, session):
        r = session.get(f"{API}/logs?limit=20", timeout=10)
        assert r.status_code == 200
        logs = r.json()
        assert isinstance(logs, list) and len(logs) > 0
        assert all("ts" in l and "message" in l for l in logs)


# ------------------------------------------------------------- config edit roundtrip
class TestConfigEditRoundtrip:
    def test_targets_put_and_persist(self, session):
        r = session.get(f"{API}/config/targets", timeout=10)
        assert r.status_code == 200
        data = r.json()
        # PUT it back as-is
        r2 = session.put(f"{API}/config/targets", json={"data": data}, timeout=15)
        assert r2.status_code == 200
        assert r2.json()["ok"] is True
        # GET back
        r3 = session.get(f"{API}/config/targets", timeout=10)
        assert r3.status_code == 200
        # structural equality (list of rows -> same length and row key set)
        got = r3.json()
        assert isinstance(got, list) and len(got) == len(data)
        assert set(got[0].keys()) == set(data[0].keys())

    def test_put_unknown_section(self, session):
        r = session.put(f"{API}/config/__nope__", json={"data": {}}, timeout=10)
        assert r.status_code == 404


# ------------------------------------------------------------- cleanup
class TestCleanup:
    def test_zz_erase_all(self, session):
        r = session.delete(f"{API}/project", timeout=15)
        assert r.status_code == 200
        st = session.get(f"{API}/state", timeout=10).json()
        assert st["project"] is None
        assert st["total_events"] == 0
        assert st["has_rating"] is False


def test_25_index_matches_excel_recomputed_data():
    """Definitive parity proof: Python's per-SDV driv index matches the value
    recomputed from the filled Excel workbook's OWN 'Indice occurrencé' column
    for all 36 scored SDVs (within 0.3). Several Excel J5 display cells are stale
    (cached, never recalculated); Python matches Excel's underlying live data,
    not the stale display. Skips gracefully if the reference workbook isn't
    present in the test environment."""
    import os, sys, uuid
    xlsx = "/mnt/user-data/uploads/25MY_2_3L_T_Ford_Bronco_P_Normal_V615_052926_ASDCA.xlsx"
    filled = "/tmp/bronco_filled.xlsm"
    if not (os.path.exists(xlsx) and os.path.exists(filled)):
        import pytest
        pytest.skip("reference workbooks not available")
    sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "backend"))
    from engine import config_loader, importer, scoring
    from engine.classifier import classify_event
    from openpyxl import load_workbook
    ss = config_loader.load_seed_sections()
    cfg = {s["key"]: s["data"] for s in ss} if isinstance(ss, list) else ss
    canon = config_loader.build_canon_map(cfg["structure"], cfg["catalog"])
    project = {"name_code": "B", "fuel": "Gasoline", "mode": "AUTO",
               "eval_mode": "Full", "gearbox": "AT", "gears": "AT",
               "number_of_gears": "10", "odriv_milestone": "QG6",
               "priority": "PREMIUM", "version": "4.6",
               "target_vehicle": "P8 MHEV MDL2", "software_milestone": "V615",
               "area": "North America", "droopy": "X"}
    ev_raw, _ = importer.parse_trie(open(xlsx, "rb").read())
    events = [{"id": str(uuid.uuid4()),
               "sdv": config_loader.canon_name(classify_event(ch, cfg["definitions"]), canon),
               "channels": ch}
              for ch in ev_raw
              if config_loader.canon_name(classify_event(ch, cfg["definitions"]), canon)]
    tl = config_loader.build_targets_lookup(cfg["targets"], project)
    _, sdv_results, _ = scoring.calculate_rating(project, events, cfg, tl, canon)
    py = {s["name"].strip().lower(): s for s in sdv_results if s.get("driv")}
    wb = load_workbook(filled, data_only=True)
    sheets = {s.lower(): s for s in wb.sheetnames}

    def recompute(sheet):
        ws = wb[sheet]
        hdr = None
        for c in range(1, 80):
            if ws.cell(6, c).value == "Indice occurrencé":
                hdr = c
                break
        if not hdr:
            return None
        last = 0
        for r in range(7, 2000):
            if ws.cell(r, 13).value not in (None, ""):
                last = r
            elif r > 7 and ws.cell(r, 13).value in (None, ""):
                break
        ios = [ws.cell(r, hdr).value for r in range(7, last + 1)
               if isinstance(ws.cell(r, hdr).value, (int, float))]
        if not ios:
            return None
        return round(100 * (1 + sum(ios) / len(ios)) ** 3, 1)

    match = tot = 0
    for k, p in py.items():
        sh = sheets.get(k)
        if not sh:
            continue
        rc = recompute(sh)
        if rc is None:
            continue
        tot += 1
        pidx = p["driv"]["index"]
        if pidx is not None and abs(pidx - rc) < 0.3:
            match += 1
    assert tot >= 30, f"expected >=30 comparable SDVs, got {tot}"
    assert match == tot, f"only {match}/{tot} SDVs match Excel's recomputed data"


def test_26_saved_projects_database(session):
    """Saved-projects database: save the processed project, list it, load it
    back after erasing, and delete it — mirroring the Excel Access DB flow."""
    session.post(f"{API}/project/new", json={
        "name_code": "DBTEST", "fuel": "Gasoline", "gearbox": "AT", "gears": "AT",
        "number_of_gears": "10", "mode": "AUTO", "eval_mode": "Full",
        "software_milestone": "QG6", "odriv_milestone": "QG6",
        "priority": "PREMIUM", "version": "4.6", "area": "North America",
        "target_vehicle": "P8 MHEV MDL2"}, timeout=30)
    session.post(f"{API}/import/demo", timeout=60)
    session.post(f"{API}/rating/calculate", timeout=120)

    r = session.post(f"{API}/db/save", timeout=60)
    assert r.status_code == 200, r.text
    saved = r.json()
    assert saved["ok"] and saved["id"]
    assert "DBTEST" in saved["uniquename"]
    pid = saved["id"]

    r = session.get(f"{API}/db/projects", timeout=30)
    assert r.status_code == 200
    lst = r.json()
    assert any(p["id"] == pid and p["code"] == "DBTEST" for p in lst["projects"])
    rec = next(p for p in lst["projects"] if p["id"] == pid)
    assert rec["event_count"] > 0
    assert rec["driv_index"] is not None

    session.delete(f"{API}/project", timeout=30)
    r = session.post(f"{API}/db/projects/{pid}/load", timeout=60)
    assert r.status_code == 200, r.text
    loaded = r.json()
    assert loaded["ok"] and loaded["project"]["name_code"] == "DBTEST"
    assert loaded["event_count"] > 0

    st = session.get(f"{API}/state", timeout=30).json()
    assert st["project"]["name_code"] == "DBTEST"
    assert st["total_events"] > 0

    r = session.delete(f"{API}/db/projects/{pid}", timeout=30)
    assert r.status_code == 200
    lst2 = session.get(f"{API}/db/projects", timeout=30).json()
    assert not any(p["id"] == pid for p in lst2["projects"])


def test_27_accdb_shared_database_read(session):
    """Shared Access .accdb: status reports availability; if a reference .accdb
    is present, list its projects and import one (read path, no Java needed)."""
    import os
    r = session.get(f"{API}/accdb/status", timeout=30)
    assert r.status_code == 200
    st = r.json()
    assert "read_available" in st and "write_available" in st
    accdb = "/mnt/user-data/uploads/_OdrivDB_1.accdb"
    if not os.path.exists(accdb):
        import pytest
        pytest.skip("reference .accdb not available")
    r = session.post(f"{API}/accdb/projects", json={"path": accdb}, timeout=120)
    assert r.status_code == 200, r.text
    projs = r.json()
    assert projs["total"] > 0
    first = projs["projects"][0]
    assert "ID" in first and "code" in first
    # import that project
    r = session.post(f"{API}/accdb/import",
                     json={"path": accdb, "project_id": first["ID"]}, timeout=180)
    assert r.status_code == 200, r.text
    imp = r.json()
    assert imp["ok"] and imp["event_count"] > 0
    assert imp["project"]["name_code"] == first["code"]


def test_28_saved_project_as_target(session):
    """Saved project selectable as comparison Target (like the Excel target
    dropdown): set shared DB path, list available targets, save a project to the
    tool DB, then apply it as the target for the current project."""
    import os
    # shared .accdb path setting persists (cfg-tab equivalent)
    r = session.put(f"{API}/settings",
                    json={"accdb_path": "/tmp/nonexistent.accdb"}, timeout=30)
    assert r.status_code == 200 and r.json()["accdb_path"].endswith(".accdb")

    # build + save a project to the tool DB to use as a target
    session.post(f"{API}/project/new", json={
        "name_code": "TGT", "fuel": "Gasoline", "gearbox": "AT", "gears": "AT",
        "number_of_gears": "10", "mode": "AUTO", "eval_mode": "Full",
        "software_milestone": "QG6", "odriv_milestone": "QG6",
        "priority": "PREMIUM", "version": "4.6", "area": "North America",
        "target_vehicle": "P8 MHEV MDL2"}, timeout=30)
    session.post(f"{API}/import/demo", timeout=60)
    session.post(f"{API}/rating/calculate", timeout=120)
    tid = session.post(f"{API}/db/save", timeout=60).json()["id"]

    # available targets include the saved project
    r = session.get(f"{API}/targets/available", timeout=60)
    assert r.status_code == 200
    avail = r.json()
    assert any(t["id"] == tid for t in avail["tool"])

    # new current project, then apply the saved one as target
    session.post(f"{API}/project/new", json={
        "name_code": "CUR", "fuel": "Gasoline", "gearbox": "AT", "gears": "AT",
        "number_of_gears": "10", "mode": "AUTO", "eval_mode": "Full",
        "software_milestone": "QG6", "odriv_milestone": "QG6",
        "priority": "PREMIUM", "version": "4.6", "area": "North America",
        "target_vehicle": "P8 MHEV MDL2"}, timeout=30)
    session.post(f"{API}/import/demo", timeout=60)
    session.post(f"{API}/rating/calculate", timeout=120)
    r = session.post(f"{API}/targets/apply",
                     json={"source": "tool", "id": tid, "label": "TGT"}, timeout=60)
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["ok"] and res["vehicle"] == "TGT" and res["rows"] > 0
    # the active project's target_vehicle is now the saved project
    st = session.get(f"{API}/state", timeout=30).json()
    assert st["project"]["target_vehicle"] == "TGT"


def test_29_accdb_folder_path(session):
    """CFG-style FOLDER path: the tool must discover _OdrivDB_*.accdb shards in a
    folder (like the Excel CFG tab) and merge projects across all of them."""
    import os
    folder = "/tmp/odriv_db_folder"
    have = os.path.isdir(folder) and any(
        f.startswith("_OdrivDB") and f.endswith(".accdb")
        for f in (os.listdir(folder) if os.path.isdir(folder) else []))
    # check-path always works (no DB needed) — validates folder handling
    r = session.post(f"{API}/accdb/check-path", json={"path": folder}, timeout=30)
    assert r.status_code == 200
    info = r.json()
    if not have:
        import pytest
        pytest.skip("reference _OdrivDB shards folder not available")
    assert info["ok"] and info["count"] >= 1
    assert all(n.endswith(".accdb") for n in info["found"])
    # listing from the FOLDER returns merged projects with shard tags
    r = session.post(f"{API}/accdb/projects", json={"path": folder}, timeout=180)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["total"] >= 1
    assert len(d["shards"]) == info["count"]
    assert any(p.get("db_file", "").endswith(".accdb") for p in d["projects"])


def test_30_accdb_year_subfolder_layout(session):
    """Real Excel layout: catalog _OdrivDB.accdb in the folder + event shards in
    a <year> subfolder. Importing must pull EVENTS from the year subfolder (the
    '0 events' bug), and Calculate must then work."""
    import os
    folder = "/tmp/odriv_real_layout"
    have = (os.path.isdir(folder)
            and os.path.exists(os.path.join(folder, "_OdrivDB.accdb"))
            and os.path.isdir(os.path.join(folder, "2024")))
    # check-path always returns 200 (folder handling); validates year detection
    r = session.post(f"{API}/accdb/check-path",
                     json={"path": folder, "year": "2024"}, timeout=30)
    assert r.status_code == 200
    if not have:
        import pytest
        pytest.skip("real-layout reference folder not available")
    info = r.json()
    assert info["ok"] and info["count"] >= 2          # catalog + >=1 shard
    assert "2024" in (info.get("year_folders") or [])
    # import a known project and confirm it has EVENTS (not zero)
    projs = session.post(f"{API}/accdb/projects",
                         json={"path": folder, "year": "2024"}, timeout=180).json()
    assert projs["total"] > 0
    pid = projs["projects"][0]["ID"]
    r = session.post(f"{API}/accdb/import",
                     json={"path": folder, "project_id": pid, "year": "2024"},
                     timeout=180)
    assert r.status_code == 200, r.text
    imp = r.json()
    assert imp["event_count"] > 0, "import must read events from the year subfolder"
    # Calculate must work now that events exist
    st = session.get(f"{API}/state", timeout=30).json()
    assert st["total_events"] > 0
    r = session.post(f"{API}/rating/calculate", timeout=120)
    assert r.status_code == 200


def test_31_multi_target_compare(session):
    """Multiple vehicles selectable as comparison targets: save two projects to
    the tool DB, apply BOTH via /targets/apply-multi, and confirm the rating
    response exposes a comparison column per vehicle."""
    ids = []
    for nm in ("MTA", "MTB"):
        session.post(f"{API}/project/new", json={
            "name_code": nm, "fuel": "Gasoline", "gearbox": "AT", "gears": "AT",
            "number_of_gears": "10", "mode": "AUTO", "eval_mode": "Full",
            "software_milestone": "QG6", "odriv_milestone": "QG6",
            "priority": "PREMIUM", "version": "4.6", "area": "North America",
            "target_vehicle": "P8 MHEV MDL2"}, timeout=30)
        session.post(f"{API}/import/demo", timeout=60)
        session.post(f"{API}/rating/calculate", timeout=120)
        ids.append(session.post(f"{API}/db/save", timeout=60).json()["id"])
    # current project
    session.post(f"{API}/project/new", json={
        "name_code": "MTCUR", "fuel": "Gasoline", "gearbox": "AT", "gears": "AT",
        "number_of_gears": "10", "mode": "AUTO", "eval_mode": "Full",
        "software_milestone": "QG6", "odriv_milestone": "QG6",
        "priority": "PREMIUM", "version": "4.6", "area": "North America",
        "target_vehicle": "P8 MHEV MDL2"}, timeout=30)
    session.post(f"{API}/import/demo", timeout=60)
    session.post(f"{API}/rating/calculate", timeout=120)
    # apply both as comparison targets
    r = session.post(f"{API}/targets/apply-multi", json={"targets": [
        {"source": "tool", "id": ids[0], "label": "MTA"},
        {"source": "tool", "id": ids[1], "label": "MTB"}]}, timeout=120)
    assert r.status_code == 200, r.text
    res = r.json()
    assert res["applied"] == 2 and set(res["vehicles"]) == {"MTA", "MTB"}
    # rating exposes both comparison columns
    rt = session.get(f"{API}/rating", timeout=30).json()
    assert set(rt.get("targets", [])) >= {"MTA", "MTB"}
    comps = rt.get("comparisons", {})
    assert "MTA" in comps and "MTB" in comps
    # each comparison has per-SDV driv indices
    assert any(v.get("driv") is not None for v in comps["MTA"].values())


def test_32_pending_targets_deferred(session):
    """Target vehicles selected in New Project are recorded instantly via
    /targets/pending (no event read). Calculate scores ONLY the current vehicle
    and leaves comparisons empty. The selected targets are read + scored only
    when 'Set as target' (/rating/set-as-target) is pressed."""
    ids = []
    for nm in ("PendA", "PendB"):
        session.post(f"{API}/project/new", json={
            "name_code": nm, "fuel": "Gasoline", "gearbox": "AT", "gears": "AT",
            "number_of_gears": "10", "mode": "AUTO", "eval_mode": "Full",
            "software_milestone": "QG6", "odriv_milestone": "QG6",
            "priority": "PREMIUM", "version": "4.6", "area": "North America",
            "target_vehicle": "P8 MHEV MDL2"}, timeout=30)
        session.post(f"{API}/import/demo", timeout=60)
        session.post(f"{API}/rating/calculate", timeout=120)
        ids.append(session.post(f"{API}/db/save", timeout=60).json()["id"])
    # current project
    session.post(f"{API}/project/new", json={
        "name_code": "PendCur", "fuel": "Gasoline", "gearbox": "AT", "gears": "AT",
        "number_of_gears": "10", "mode": "AUTO", "eval_mode": "Full",
        "software_milestone": "QG6", "odriv_milestone": "QG6",
        "priority": "PREMIUM", "version": "4.6", "area": "North America",
        "target_vehicle": "P8 MHEV MDL2"}, timeout=30)
    session.post(f"{API}/import/demo", timeout=60)
    # set pending — must be fast and must NOT populate comparisons yet
    import time
    t = time.time()
    r = session.post(f"{API}/targets/pending", json={
        "primary": {"source": "tool", "id": ids[0], "label": "PendA"},
        "extras": [{"source": "tool", "id": ids[1], "label": "PendB"}]}, timeout=30)
    assert r.status_code == 200, r.text
    assert r.json()["pending"] == 2
    assert (time.time() - t) < 5      # instant (no event read)
    # before set-as-target, comparisons should be empty (deferred)
    rt = session.get(f"{API}/rating", timeout=30).json()
    assert not rt.get("comparisons")
    # Calculate scores ONLY the current vehicle -> still no comparisons
    session.post(f"{API}/rating/calculate", timeout=180)
    rt = session.get(f"{API}/rating", timeout=30).json()
    assert not rt.get("comparisons"), "Calculate must not resolve target vehicles"
    # 'Set as target' reads + scores the selected target vehicles
    r = session.post(f"{API}/rating/set-as-target", timeout=120)
    assert r.status_code == 200, r.text
    rt = session.get(f"{API}/rating", timeout=30).json()
    assert "PendA" in rt.get("comparisons", {}) and "PendB" in rt.get("comparisons", {})


def test_33_calculate_progress(session):
    """The background calculate exposes a progress percentage that reaches 100%
    and a done flag, and produces the same scored result."""
    session.post(f"{API}/project/new", json={
        "name_code": "ProgT", "fuel": "Gasoline", "gearbox": "AT", "gears": "AT",
        "number_of_gears": "10", "mode": "AUTO", "eval_mode": "Full",
        "software_milestone": "QG6", "odriv_milestone": "QG6",
        "priority": "PREMIUM", "version": "4.6", "area": "North America",
        "target_vehicle": "P8 MHEV MDL2"}, timeout=30)
    session.post(f"{API}/import/demo", timeout=60)
    r = session.post(f"{API}/rating/calculate/start", timeout=30)
    assert r.status_code == 200
    # poll until done
    import time
    done = False
    for _ in range(120):
        time.sleep(0.3)
        p = session.get(f"{API}/rating/progress", timeout=30).json()
        assert 0 <= p["percent"] <= 100
        if p["done"]:
            done = True
            assert p["error"] is None
            break
    assert done, "calculation did not finish"
    # rating is populated afterwards
    rt = session.get(f"{API}/rating", timeout=30).json()
    assert rt.get("rows")
