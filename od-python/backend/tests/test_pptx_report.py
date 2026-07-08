"""Unit tests for PowerPoint report generation."""
import os
import tempfile

import pytest
from pptx import Presentation

from engine.pptx_report import build_pptx


@pytest.fixture
def sample_report_data():
    return {
        "project": {
            "name_code": "TEST_PROJECT",
            "mode": "AUTO",
            "fuel": "Diesel",
            "gears": "8AT",
            "odriv_milestone": "MS3",
            "version": "29",
            "area": "EU",
            "target_vehicle": "REF",
        },
        "global": {
            "driv": {
                "verdict": "Medium Risk",
                "verdict_pred": "Low Risk",
                "index": 72.5,
                "rate_low": 0.12,
            },
            "dyn": {
                "verdict": "Low Risk",
                "verdict_pred": "Low Risk",
                "index": 81.0,
                "rate_low": 0.05,
            },
        },
        "sdv_results": [
            {
                "name": "Drive Away Creep Eng On",
                "order": 1,
                "n_events": 5,
                "driv": {
                    "index": 70.2,
                    "target_index": 75.0,
                    "status": {"1": "RED", "2": "GREEN", "3": "GREEN"},
                    "status_pred": {"1": "ORANGE", "2": "GREEN", "3": "GREEN"},
                    "counts": {"RED_P1": 1, "GREEN_P1": 4},
                    "lowest_event": "SubEvent-1",
                },
                "dyn": {
                    "index": 78.0,
                    "target_index": 80.0,
                    "status": {"1": "GREEN", "2": "GREEN", "3": "GREEN"},
                    "status_pred": {"1": "GREEN", "2": "GREEN", "3": "GREEN"},
                    "counts": {"GREEN_P1": 5},
                    "lowest_event": None,
                },
            }
        ],
        "charts": {"Drive Away Creep Eng On": {"png": None}},
        "events_by_sdv": {
            "Drive Away Creep Eng On": [
                {
                    "channels": [{"name": "Sub Event Name", "value": "E1"}],
                    "driv": {
                        "criticity": 1,
                        "priority": 1,
                        "color": "RED",
                        "indice_occ": -2.5,
                    },
                }
            ]
        },
        "doc_versions": [{"name": "ODRIV", "version": "v29 test"}],
    }


def test_build_pptx_creates_valid_file(sample_report_data):
    fd, out = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    try:
        build_pptx(
            sample_report_data,
            out,
            report_fields={"sender_name": "Unit Test", "domain": "DRIVABILITY-DYNAMISM"},
        )
        assert os.path.getsize(out) > 10000
        prs = Presentation(out)
        assert len(prs.slides) >= 8
        titles = []
        for slide in prs.slides:
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    titles.append(sh.text_frame.text.strip().split("\n")[0])
                    break
        joined = " | ".join(titles).upper()
        assert "OBJECTIVE VALIDATION REPORT" in joined or "REPORT METADATA" in joined
        assert "DRIVE AWAY CREEP ENG ON" in joined
        assert "ANNEXES" in joined
    finally:
        if os.path.exists(out):
            os.unlink(out)
