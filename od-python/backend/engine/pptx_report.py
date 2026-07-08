"""PowerPoint (.pptx) ODrive report — same structure as the Word report / Excel tool."""
import io
from datetime import datetime

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from engine.report_tables import (
    SYNTHESIS_PARTS,
    _global_risk_png,
    _has_dyn_section,
    _png_bytes,
    _priority_table_png,
    _project_home_png,
    _scorecard_png_chunks,
    _sdv_synthesis_png,
)
from engine.reports import TEAL, _fmt_doc_version

NAVY = (0x1E, 0x23, 0x36)


def _slide_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.4), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*TEAL)


def _add_text_block(slide, left, top, width, height, lines, font_size=14, bold_first=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line
        para.font.size = Pt(font_size)
        if bold_first and i == 0:
            para.font.bold = True


def _add_table_slide(prs, blank, title, table_rows):
    """table_rows: list of (label, value) pairs."""
    slide = prs.slides.add_slide(blank)
    _slide_title(slide, title)
    n = len(table_rows)
    shape = slide.shapes.add_table(n, 2, Inches(0.5), Inches(1.1), Inches(12.3), Inches(0.35 * n))
    table = shape.table
    for r, (label, value) in enumerate(table_rows):
        table.cell(r, 0).text = str(label)
        table.cell(r, 1).text = str(value)
        for c in (0, 1):
            for para in table.cell(r, c).text_frame.paragraphs:
                para.font.size = Pt(11)
        table.cell(r, 0).text_frame.paragraphs[0].font.bold = True
    return slide


def _add_image_slide(prs, blank, title, png_buf, subtitle=None, img_width=12.0, top=1.05):
    slide = prs.slides.add_slide(blank)
    _slide_title(slide, title)
    y = top
    if subtitle:
        _add_text_block(slide, Inches(0.5), Inches(y), Inches(12), Inches(0.35), [subtitle], font_size=12, bold_first=True)
        y += 0.45
    raw = _png_bytes(png_buf)
    if raw:
        slide.shapes.add_picture(io.BytesIO(raw), Inches(0.5), Inches(y), width=Inches(img_width))
    return slide


def _cover_metadata_rows(project, fields):
    p = project or {}
    f = fields or {}
    return [
        ("Topic", "OBJECTIVE VALIDATION REPORT — SHIFT QUALITY/ENGINE/POWERTRAIN"),
        ("Domain", f.get("domain") or "DRIVABILITY-DYNAMISM"),
        ("Project", f.get("project") or p.get("name_code") or "-"),
        ("Standard", f.get("standard") or "n/a"),
        ("Stage", f.get("stage") or "n/a"),
        ("Goal", f.get("goal") or "n/a"),
    ]


def _sender_lines(fields):
    f = fields or {}
    return [
        "SENDER(S)",
        "  From: %s" % (f.get("sender_name") or "n/a"),
        "  Department: %s" % (f.get("sender_dept") or "n/a"),
        "  Tél.: %s    Mail to: %s" % (f.get("sender_tel") or "n/a", f.get("sender_email") or "n/a"),
        "",
        "RECIPIENT(S): %s" % (f.get("recipients") or "For Information"),
        "",
        "Place, %s" % (f.get("site") or "n/a"),
        datetime.now().strftime("%m/%d/%Y"),
    ]


def _synthesis_lines(fields):
    f = fields or {}
    lines = ["SYNTHESIS:"]
    syn = (f.get("synthesis") or "").strip()
    if syn:
        lines.extend(syn.splitlines())
    else:
        lines.append("(Add synthesis narrative in report header fields)")
    dyn = (f.get("dynamism_synthesis") or "").strip()
    if dyn:
        lines.extend(["", "DYNAMISM SYNTHESIS:", dyn])
    return lines


def _test_conditions_lines(fields):
    f = fields or {}
    return [
        "Test conditions",
        "- Location and date of test: %s" % (f.get("loc_date_test") or "n/a"),
        "- Climate condition: %s" % (f.get("climate") or "n/a"),
        "- A/C status: %s" % (f.get("ac_status") or "n/a"),
        "- Vehicle options: %s" % (f.get("vehicle_options") or "n/a"),
    ]


def _doc_versions_rows(doc_versions):
    links = {
        "self-learning tool": "http://docinfogroupe.inetpsa.com/ead/doc/ref.01471_17_01042/v.vc/fiche",
        "event collector": "http://docinfogroupe.inetpsa.com/ead/doc/ref.01470_16_00083/v.vc/fiche",
        "objective evaluation": "http://docinfogroupe.inetpsa.com/ead/doc/ref.01470_15_00987/v.vc/fiche",
        "experiment": "http://docinfogroupe.inetpsa.com/ead/doc/ref.01472_17_03797/v.vc/fiche",
        "odriv": "http://docinfogroupe.inetpsa.com/ead/doc/ref.01470_16_00345/v.vc/fiche",
    }
    rows = []
    for key, url in links.items():
        ver = ""
        for dv in doc_versions or []:
            if isinstance(dv, dict) and key in dv.get("name", "").lower():
                ver = dv.get("version") or ""
                break
        if key == "odriv" and not ver and doc_versions:
            ver = _fmt_doc_version(doc_versions[-1])
        rows.append((key.title(), "%s | %s" % (url, ver or "n/a")))
    return rows


def _add_sdv_slides(prs, blank, data, events_by_sdv, charts):
    for j, sdv in enumerate(data.get("sdv_results") or []):
        name = sdv.get("name") or "SDV"
        events = events_by_sdv.get(name) or []
        chart_png = (charts.get(name) or {}).get("png")
        for t, part_key, part_label in ((1, "driv", "DRIVABILITY"), (2, "dyn", "RESPONSIVENESS")):
            if t == 2 and not _has_dyn_section(sdv):
                continue
            if t == 2 and not (sdv.get("dyn") or {}):
                continue
            heading = "2.%d.%d %s %s" % (j + 1, t, name.upper(), part_label)
            slide = prs.slides.add_slide(blank)
            _slide_title(slide, heading)
            y = 1.0
            sections = [
                (SYNTHESIS_PARTS[0], _sdv_synthesis_png(sdv, part_key)),
                (SYNTHESIS_PARTS[1], chart_png if chart_png else None),
            ]
            hi = _priority_table_png(events, part_key, "high")
            lo = _priority_table_png(events, part_key, "low")
            if hi:
                sections.append((SYNTHESIS_PARTS[3], hi))
            if lo:
                sections.append((SYNTHESIS_PARTS[4], lo))
            for subtitle, png in sections:
                if not png:
                    continue
                _add_text_block(slide, Inches(0.5), Inches(y), Inches(12), Inches(0.3),
                                [subtitle], font_size=12, bold_first=True)
                y += 0.38
                raw = _png_bytes(png)
                if raw:
                    pic = slide.shapes.add_picture(io.BytesIO(raw), Inches(0.5), Inches(y), width=Inches(6.2))
                    y += pic.height.inches + 0.2
                if y > 6.8:
                    slide = prs.slides.add_slide(blank)
                    _slide_title(slide, heading + " (continued)")
                    y = 1.0


def build_pptx(data, out_path, report_fields=None):
    """Build an ODrive PowerPoint report matching the Word / Excel tool layout."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    project = data.get("project") or {}
    glob = data.get("global") or {}
    fields = dict(report_fields or {})
    if not fields.get("project"):
        fields["project"] = project.get("name_code") or ""

    # Title slide
    slide = prs.slides.add_slide(blank)
    band = slide.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(12), Inches(1.2))
    p = band.text_frame.paragraphs[0]
    p.text = "OBJECTIVE VALIDATION REPORT"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*NAVY)
    sub = slide.shapes.add_textbox(Inches(0.6), Inches(3.2), Inches(12), Inches(1.0))
    sub.text_frame.paragraphs[0].text = "SHIFT QUALITY / ENGINE / POWERTRAIN — ODRIV"
    sub.text_frame.paragraphs[0].font.size = Pt(18)

    _add_table_slide(prs, blank, "Report metadata", _cover_metadata_rows(project, fields))

    slide = prs.slides.add_slide(blank)
    _slide_title(slide, "Sender / Recipients")
    _add_text_block(slide, Inches(0.5), Inches(1.1), Inches(12), Inches(2.5), _sender_lines(fields), font_size=13)

    slide = prs.slides.add_slide(blank)
    _slide_title(slide, "SYNTHESIS")
    _add_text_block(slide, Inches(0.5), Inches(1.1), Inches(12), Inches(5.5), _synthesis_lines(fields), font_size=12)

    _add_table_slide(prs, blank, "Document versions (DocInfo)", _doc_versions_rows(data.get("doc_versions") or []))

    slide = prs.slides.add_slide(blank)
    _slide_title(slide, "SYNTHESIS — Test conditions")
    _add_text_block(slide, Inches(0.5), Inches(1.1), Inches(12), Inches(2.0), _test_conditions_lines(fields), font_size=12)

    home_png = _project_home_png(project, glob)
    if home_png:
        _add_image_slide(prs, blank, "Project summary", home_png)

    risk_png = _global_risk_png(glob)
    if risk_png:
        _add_image_slide(prs, blank, "Risk assessment for customer complaints", risk_png)

    for title, png in _scorecard_png_chunks(data):
        _add_image_slide(prs, blank, title, png)

    slide = prs.slides.add_slide(blank)
    _slide_title(slide, "Objective results")

    _add_sdv_slides(prs, blank, data, data.get("events_by_sdv") or {}, data.get("charts") or {})

    slide = prs.slides.add_slide(blank)
    _slide_title(slide, "ANNEXES")
    _add_text_block(slide, Inches(0.5), Inches(1.1), Inches(12), Inches(4.5), [
        "Other remarks",
        "Insert a MDA view of low points that AVL Drive does not objectively assess.",
        "",
        "Problems",
        "List all problems encountered during the evaluation.",
        "",
        "OIL docinfo link",
        "Please insert the docinfo link concerning this powertrain.",
    ], font_size=12)

    prs.save(out_path)
    return out_path
